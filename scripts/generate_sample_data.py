"""生成电商商品样例知识库文档（用于演示与答辩）。

输出：scripts/sample_data/
  - 手机 5 款  → xlsx（参数表 / 价格政策 / 售后条款 三个 sheet）
  - 家电 5 款  → docx（段落 + 表格）
  - 数码 5 款  → txt
  - 其他 5 款  → md（服饰/食品/家居/运动/美妆）
  - 平台售后总则 1 份 → pdf（reportlab 内置中文字体）

用法：cd backend && venv/Scripts/python ../scripts/generate_sample_data.py
"""
import os
import sys
from pathlib import Path

# 允许从任意目录运行
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "backend"))

OUT_DIR = Path(__file__).resolve().parent / "sample_data"

# ==================== 数据 ====================
PHONES = [
    {
        "name": "星耀 X90 Pro 5G",
        "params": [
            ("屏幕", "6.78 英寸 AMOLED，1.5K 分辨率，120Hz 刷新率"),
            ("处理器", "天玑 9400+ 旗舰芯片"),
            ("运行内存", "12GB / 16GB LPDDR5X"),
            ("机身存储", "256GB / 512GB / 1TB"),
            ("电池容量", "5800mAh"),
            ("充电", "120W 有线快充，50W 无线快充"),
            ("摄像头", "后置 5000 万像素主摄 + 5000 万超广角 + 6400 万潜望长焦"),
            ("前置摄像头", "3200 万像素"),
            ("系统", "星耀 OS 5.0（基于 Android 15）"),
            ("重量", "208g"),
            ("颜色", "曜石黑 / 星尘白 / 远山青"),
            ("其他特性", "支持 NFC、红外遥控、IP68 防尘防水、双扬声器"),
        ],
        "price": "首发价 4299 元起（12GB+256GB）。平台补贴价 3999 元，学生认证再减 100 元。支持 12 期免息分期，赠原装 120W 充电器。",
        "aftersale": "7 天无理由退货（激活后不支持），15 天内质量问题换新，1 年整机保修，电池健康度低于 80% 可在保修期内免费更换电池。碎屏险 99 元/年。",
    },
    {
        "name": "星耀 X90 标准版",
        "params": [
            ("屏幕", "6.67 英寸 OLED，120Hz 刷新率"),
            ("处理器", "天玑 8300 芯片"),
            ("运行内存", "8GB / 12GB"),
            ("机身存储", "128GB / 256GB"),
            ("电池容量", "5100mAh"),
            ("充电", "66W 有线快充"),
            ("摄像头", "后置 5000 万主摄 + 800 万超广角"),
            ("重量", "189g"),
            ("其他特性", "支持 NFC、立体声双扬声器"),
        ],
        "price": "首发价 2499 元起（8GB+128GB）。平台补贴价 2299 元，支持 6 期免息分期。",
        "aftersale": "7 天无理由退货，15 天质量问题换新，1 年整机保修。",
    },
    {
        "name": "华光 MateBook 手机 60 青春版",
        "params": [
            ("屏幕", "6.5 英寸 LCD，90Hz 刷新率"),
            ("处理器", "骁龙 7s Gen2"),
            ("运行内存", "8GB"),
            ("机身存储", "128GB / 256GB"),
            ("电池容量", "6000mAh"),
            ("充电", "44W 快充"),
            ("摄像头", "后置 1 亿像素主摄"),
            ("重量", "199g"),
            ("其他特性", "3.5mm 耳机孔、侧边指纹"),
        ],
        "price": "售价 1799 元起，平台补贴价 1599 元，不支持分期。",
        "aftersale": "7 天无理由退货，15 天质量问题换新，1 年整机保修。",
    },
    {
        "name": "华光 P90 折叠屏",
        "params": [
            ("屏幕", "外屏 6.4 英寸 + 内屏 7.9 英寸，均为 120Hz LTPO"),
            ("处理器", "骁龙 8 Gen4"),
            ("运行内存", "16GB"),
            ("机身存储", "512GB / 1TB"),
            ("电池容量", "5600mAh"),
            ("充电", "66W 有线 + 40W 无线"),
            ("摄像头", "后置 5000 万四摄，支持 5 倍光学变焦"),
            ("重量", "239g"),
            ("其他特性", "IPX8 防水、卫星通信"),
        ],
        "price": "售价 9999 元起，平台补贴价 8999 元，支持 24 期免息分期，赠折叠屏专属保护壳。",
        "aftersale": "7 天无理由退货，15 天质量问题换新，2 年整机保修（含铰链），免费贴膜服务 2 次/年。",
    },
    {
        "name": "睿米 15 Ultra",
        "params": [
            ("屏幕", "6.73 英寸 2K AMOLED，120Hz"),
            ("处理器", "骁龙 8 Gen4"),
            ("运行内存", "16GB"),
            ("机身存储", "512GB / 1TB"),
            ("电池容量", "5410mAh"),
            ("充电", "90W 有线 + 80W 无线"),
            ("摄像头", "后置徕卡 5000 万四摄，1 英寸大底主摄"),
            ("重量", "226g"),
            ("其他特性", "IP68 防水、超声波指纹"),
        ],
        "price": "售价 6499 元起，平台补贴价 5999 元，支持 12 期免息分期，以旧换新至高补贴 800 元。",
        "aftersale": "7 天无理由退货，15 天质量问题换新，1 年整机保修，屏幕与主板 2 年延保服务。",
    },
]

APPLIANCES = [
    {
        "name": "云洁变频滚筒洗衣机 WX-1012",
        "params": [
            ("容量", "10kg 大容量"),
            ("电机", "直驱变频电机"),
            ("能效等级", "一级能效"),
            ("洗净比", "1.10"),
            ("特色功能", "智能投放洗衣液、除菌螨洗、空气洗、静音运行 48dB"),
            ("尺寸", "595×600×850mm"),
        ],
        "price": "售价 3299 元，平台补贴价 2799 元，支持 3 期免息，免费上门安装，以旧换新至高抵 300 元。",
        "aftersale": "整机 3 年保修，电机 10 年包修，7 天无理由退货（未安装），30 天质量问题包退，90 天包换。",
    },
    {
        "name": "云洁一级能效变频空调 KFR-35GW",
        "params": [
            ("匹数", "1.5 匹（适用 16-22㎡）"),
            ("能效等级", "一级能效（APF 5.3）"),
            ("制冷量", "3500W"),
            ("制热量", "5000W"),
            ("噪音", "18dB 静音"),
            ("特色功能", "新风系统、自清洁、智能互联、防直吹"),
        ],
        "price": "售价 2699 元，平台补贴价 2299 元，免费上门安装（含支架），高空作业费另计，支持 3 期免息。",
        "aftersale": "整机 6 年保修，压缩机 10 年包修，7 天无理由退货（未安装），30 天质量问题包退。",
    },
    {
        "name": "云洁对开门冰箱 BCD-540",
        "params": [
            ("总容积", "540L（冷藏 340L + 冷冻 200L）"),
            ("能效等级", "一级能效"),
            ("制冷方式", "风冷无霜"),
            ("特色功能", "双变频、独立变温室、净味除菌、智能互联"),
            ("噪音", "38dB"),
            ("尺寸", "910×700×1800mm"),
        ],
        "price": "售价 3999 元，平台补贴价 3499 元，免费送货入户，支持 3 期免息，以旧换新至高抵 400 元。",
        "aftersale": "整机 3 年保修，压缩机 10 年包修，7 天无理由退货（未通电），30 天质量问题包退。",
    },
    {
        "name": "云洁扫地机器人 S9 Pro",
        "params": [
            ("吸力", "8000Pa 大吸力"),
            ("导航", "激光雷达 + AI 视觉避障"),
            ("拖地", "旋转加压拖布，自动回洗"),
            ("基站", "自动集尘、自动烘干、自动补水"),
            ("续航", "180 分钟，自动回充"),
            ("特色功能", "语音控制、App 建图、地毯识别"),
        ],
        "price": "售价 4299 元，平台补贴价 3699 元，赠耗材礼包（边刷×2、拖布×2、尘袋×3），支持 6 期免息。",
        "aftersale": "整机 2 年保修，电池 1 年保修，7 天无理由退货，30 天质量问题包退，以换代修服务。",
    },
    {
        "name": "云洁即热式净水器 WQ-800",
        "params": [
            ("过滤精度", "RO 反渗透 0.0001 微米"),
            ("通量", "800G 大通量"),
            ("出水速度", "2.1L/min"),
            ("滤芯", "复合滤芯 + RO 滤芯，总净水量 6000L"),
            ("特色功能", "即滤即热、四档水温、智能 TDS 显示、App 换芯提醒"),
        ],
        "price": "售价 1999 元，平台补贴价 1699 元，免费上门安装，首年滤芯 8 折。",
        "aftersale": "整机 3 年保修，滤芯属耗材不在保修范围，7 天无理由退货（未安装）。",
    },
]

DIGITAL = [
    {
        "name": "轻羽无线降噪耳机 Air Pro 3",
        "params": [
            ("降噪深度", "48dB 自适应主动降噪"),
            ("蓝牙版本", "蓝牙 5.4"),
            ("续航", "单次 9 小时，配合充电盒 36 小时"),
            ("充电", "Type-C 快充，充电 10 分钟听歌 3 小时"),
            ("扬声器", "12mm 动圈单元，Hi-Res 金标认证"),
            ("重量", "单耳 4.6g"),
            ("特色功能", "空间音频、双设备连接、IP54 防水"),
        ],
        "price": "售价 499 元，平台补贴价 399 元，不支持分期。",
        "aftersale": "7 天无理由退货，15 天质量问题换新，1 年保修。",
    },
    {
        "name": "轻羽平板 TabPad 11",
        "params": [
            ("屏幕", "11 英寸 2.4K LCD，120Hz 刷新率"),
            ("处理器", "骁龙 8s Gen3"),
            ("内存", "8GB / 12GB"),
            ("存储", "128GB / 256GB"),
            ("电池", "8600mAh，67W 快充"),
            ("特色功能", "四扬声器、手写笔磁吸充电、键盘保护套接口"),
        ],
        "price": "售价 2499 元起，平台补贴价 2199 元起，支持 6 期免息，教育优惠赠手写笔。",
        "aftersale": "7 天无理由退货，15 天质量问题换新，1 年整机保修。",
    },
    {
        "name": "轻羽智能手表 Watch S4",
        "params": [
            ("屏幕", "1.43 英寸 AMOLED 常亮屏"),
            ("续航", "典型使用 14 天"),
            ("防水", "5ATM 防水"),
            ("健康监测", "心率、血氧、睡眠、压力、女性健康"),
            ("运动模式", "120+ 种运动模式，双频 GPS"),
            ("特色功能", "蓝牙通话、NFC 门禁公交卡、微信手表版"),
        ],
        "price": "售价 999 元，平台补贴价 849 元，支持 3 期免息。",
        "aftersale": "7 天无理由退货，15 天质量问题换新，1 年保修。",
    },
    {
        "name": "轻羽智能音箱 Echo Mini",
        "params": [
            ("扬声器", "2 英寸全频单元 + 被动辐射器"),
            ("麦克风", "4 麦克风阵列，远场拾音 5 米"),
            ("连接", "Wi-Fi 2.4G/5G、蓝牙 5.3"),
            ("特色功能", "智能家居控制中枢、儿童模式、多设备组网播放"),
        ],
        "price": "售价 199 元，平台补贴价 159 元，不支持分期。",
        "aftersale": "7 天无理由退货，15 天质量问题换新，1 年保修。",
    },
    {
        "name": "轻羽移动电源 20000mAh 快充版",
        "params": [
            ("容量", "20000mAh（74Wh，可上飞机）"),
            ("输出", "最高 65W PD 快充，可充笔记本"),
            ("接口", "2×USB-C + 1×USB-A"),
            ("输入", "65W 快充，2.5 小时充满"),
            ("特色功能", "数显电量、小电流模式（充耳机手表）"),
        ],
        "price": "售价 169 元，平台补贴价 129 元，不支持分期。",
        "aftersale": "7 天无理由退货，15 天质量问题换新，1 年保修（鼓包免费换新）。",
    },
]

OTHERS = [
    {
        "name": "云感纯棉基础款T恤",
        "params": [
            ("面料", "100% 新疆长绒棉，180g 重磅"),
            ("版型", "微宽松落肩版型"),
            ("尺码", "S / M / L / XL / XXL（标准尺码，不偏码）"),
            ("颜色", "白色 / 黑色 / 灰色 / 藏青"),
            ("工艺", "二本针领口，不易变形"),
        ],
        "price": "售价 79 元/件，第二件半价，平台补贴后两件 118.5 元。支持 7 天无理由退换（吊牌完好）。",
        "aftersale": "7 天无理由退换货，质量问题（破洞、色差严重）运费由商家承担，尺码不合适支持免费换码一次。",
    },
    {
        "name": "山萃每日坚果 30 包礼盒",
        "params": [
            ("规格", "25g×30 包（750g）"),
            ("配料", "巴旦木、腰果、核桃仁、榛子仁、蔓越莓干、蓝莓干"),
            ("保质期", "180 天"),
            ("储存", "阴凉干燥处，开封后建议 3 天内食用"),
            ("营养", "每包约 140 千卡，0 添加蔗糖"),
        ],
        "price": "售价 89 元，平台补贴价 69 元，两件 9 折，临期（保质期剩 45 天内）特价 39 元。",
        "aftersale": "食品类不支持 7 天无理由退货；漏气、发霉等质量问题拍照即可退款，无需退货。",
    },
    {
        "name": "云居记忆棉慢回弹枕",
        "params": [
            ("枕芯", "60D 慢回弹记忆棉"),
            ("尺寸", "60×40×11cm"),
            ("枕套", "可拆洗天竺棉枕套"),
            ("适用人群", "仰睡/侧睡人群，颈椎不适人群"),
            ("特色", "人体工学曲线分区，透气孔设计"),
        ],
        "price": "售价 199 元，平台补贴价 149 元，买一送一（同款枕套）。",
        "aftersale": "支持 30 天试睡，不满意全额退款（退回运费自理）；枕芯 2 年塌陷包换。",
    },
    {
        "name": "跃动碳素羽毛球拍套装",
        "params": [
            ("材质", "全碳素拍框 + 碳纤维中杆"),
            ("重量", "4U（约 83g）"),
            ("平衡点", "295mm 偏进攻"),
            ("拉线磅数", "最高 30 磅"),
            ("套装内容", "球拍×2、羽毛球×3、拍包×1、手胶×4"),
        ],
        "price": "售价 299 元/对，平台补贴价 239 元/对，支持 3 期免息。",
        "aftersale": "7 天无理由退货（未穿线），拍框断裂 90 天内免费换新，手柄缠胶免费。",
    },
    {
        "name": "花漾保湿修护精华液",
        "params": [
            ("规格", "30ml"),
            ("核心成分", "神经酰胺、烟酰胺、透明质酸钠、积雪草提取物"),
            ("适用肤质", "干性、中性、混合性肌肤；敏感肌可用"),
            ("功效", "补水保湿、修护屏障、提亮肤色"),
            ("保质期", "3 年，开封后 12 个月内用完"),
        ],
        "price": "售价 259 元，平台补贴价 199 元，买赠同款 10ml 旅行装。",
        "aftersale": "美妆类开封后不支持退换；过敏（提供照片）48 小时内可申请退款；未拆封支持 7 天无理由。",
    },
]

PLATFORM_POLICY = """慧购优选平台售后服务总则

一、退换货政策
1. 7 天无理由退货：自签收次日起 7 天内，商品完好（不影响二次销售）可申请无理由退货，运费按商品页面标注规则承担。
2. 特殊品类：食品、美妆（已开封）、内衣、定制商品不支持无理由退货，具体以商品页面为准。
3. 质量问题：签收后 15 天内出现质量问题可申请换新；保修期内按商品页面保修政策执行。

二、价格保护
1. 自下单之日起 7 天内，若同商品发生降价，可申请退还差价（限时秒杀、直播专享价除外）。
2. 平台补贴价与商家优惠券可叠加使用，补贴金额以结算页为准。

三、发票
1. 支持开具电子发票（增值税普通发票/专用发票），下单时勾选或收货后 30 天内补开。
2. 电子发票与纸质发票具有同等法律效力。

四、客服与投诉
1. 在线客服 7×24 小时服务，人工客服 9:00-24:00。
2. 对处理结果不满可发起平台介入，平台将在 48 小时内给出仲裁结果。

五、物流
1. 平台自营商品 48 小时内发货，第三方商家 72 小时内发货（预售商品除外）。
2. 支持物流全程追踪，签收前请当面验货，破损可拒收。"""


# ==================== 生成函数 ====================
def _write_xlsx(path: Path, product: dict) -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "商品参数"
    ws.append(["参数项", "参数值"])
    for k, v in product["params"]:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 60
    ws["A1"].font = ws["B1"].font = Font(bold=True)

    ws2 = wb.create_sheet("价格政策")
    ws2.append(["政策类别", "内容"])
    ws2.append(["定价与促销", product["price"]])
    ws2.append(["分期与赠品", "以商品页面实时展示为准"])

    ws3 = wb.create_sheet("售后条款")
    ws3.append(["条款类别", "内容"])
    ws3.append(["售后承诺", product["aftersale"]])

    wb.save(path)


def _write_docx(path: Path, product: dict) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(product["name"], level=1)
    doc.add_heading("商品参数", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "参数项"
    table.rows[0].cells[1].text = "参数值"
    for k, v in product["params"]:
        row = table.add_row()
        row.cells[0].text = k
        row.cells[1].text = v
    doc.add_heading("价格政策", level=2)
    doc.add_paragraph(product["price"])
    doc.add_heading("售后条款", level=2)
    doc.add_paragraph(product["aftersale"])
    doc.save(path)


def _write_txt(path: Path, product: dict) -> None:
    lines = [f"商品名称：{product['name']}", "", "【商品参数】"]
    lines += [f"- {k}：{v}" for k, v in product["params"]]
    lines += ["", "【价格政策】", product["price"], "", "【售后条款】", product["aftersale"]]
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_md(path: Path, product: dict) -> None:
    lines = [f"# {product['name']}", "", "## 商品参数", ""]
    lines += [f"- **{k}**：{v}" for k, v in product["params"]]
    lines += ["", "## 价格政策", "", product["price"], "", "## 售后条款", "", product["aftersale"]]
    path.write_text("\n".join(lines), encoding="utf-8")


def _write_pdf(path: Path, text: str) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    style = ParagraphStyle("cn", fontName="STSong-Light", fontSize=11, leading=18)
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    story = [
        Paragraph(f"<p>{p}</p>", style)
        for p in text.split("\n")
    ]
    doc.build(story)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    count = 0
    for i, p in enumerate(PHONES):
        _write_xlsx(OUT_DIR / f"手机_{i + 1}_{p['name']}_参数价格售后.xlsx", p)
        count += 1
    for i, p in enumerate(APPLIANCES):
        _write_docx(OUT_DIR / f"家电_{i + 1}_{p['name']}.docx", p)
        count += 1
    for i, p in enumerate(DIGITAL):
        _write_txt(OUT_DIR / f"数码_{i + 1}_{p['name']}.txt", p)
        count += 1
    for i, p in enumerate(OTHERS):
        _write_md(OUT_DIR / f"其他_{i + 1}_{p['name']}.md", p)
        count += 1
    _write_pdf(OUT_DIR / "平台售后服务总则.pdf", PLATFORM_POLICY)
    count += 1
    print(f"生成完成：{count} 个文件 → {OUT_DIR}")


if __name__ == "__main__":
    main()
