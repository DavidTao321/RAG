"""文档解析器测试：txt 编码探测、docx 段落+表格、excel 表头/合并单元格、pdf 主链路与回退。

样例文件全部在 tmp_path 动态生成（不提交二进制文件到仓库）；pdf 两条链路用 monkeypatch 打桩，离线确定。
"""
import pandas as pd
import pytest
from docx import Document

from app.services.parsers import PARSERS, parse_file
from app.services.parsers.docx_parser import _table_to_lines, parse_docx
from app.services.parsers.excel_parser import parse_excel
from app.services.parsers.pdf_parser import parse_pdf
from app.services.parsers.text_parser import parse_text


# ===== txt =====

def test_parse_text_utf8(tmp_path):
    """utf-8 文本正常解析。"""
    p = tmp_path / "a.txt"
    p.write_text("充电器：氮化镓 65W", encoding="utf-8")
    assert parse_text(str(p)) == "充电器：氮化镓 65W"


def test_parse_text_gbk(tmp_path):
    """gbk 编码文本靠 charset-normalizer 探测解析（中文 Windows 常见编码）。

    注意：过短的中文片段在 GBK/EUC-KR 间有字节歧义会被误判，用足够长的文本保证探测确定。
    """
    content = "星耀X90手机详细参数表：屏幕6.78英寸、处理器天玑9400+、电池容量5800mAh、充电120W、后置摄像头5000万像素、操作系统、颜色、价格与售后政策说明"
    p = tmp_path / "b.txt"
    p.write_bytes(content.encode("gbk"))
    assert parse_text(str(p)) == content


# ===== docx =====

def _make_docx(path):
    doc = Document()
    doc.add_paragraph("星耀 X90 手机参数")
    doc.add_paragraph("")  # 空段落应被跳过
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "参数"
    table.rows[0].cells[1].text = "数值"
    table.rows[1].cells[0].text = "电池"
    table.rows[1].cells[1].text = "5800mAh"
    doc.save(str(path))


def test_parse_docx_paragraphs_and_table(tmp_path):
    """段落文本 + 表格行转「名称：数值」键值对（每行两两配对），空段落被过滤。"""
    p = tmp_path / "a.docx"
    _make_docx(p)
    text = parse_docx(str(p))
    assert text == "星耀 X90 手机参数\n参数：数值\n电池：5800mAh"


def test_table_to_lines_odd_cells(tmp_path):
    """奇数单元格行：前两两配对后剩余单元格单独追加。"""
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    cells = table.rows[0].cells
    cells[0].text = "A"
    cells[1].text = "B"
    cells[2].text = "C"
    assert _table_to_lines(table) == ["A：B；C"]


def test_table_to_lines_newline_in_cell_replaced(tmp_path):
    """单元格内换行被替换为空格，避免破坏行结构。"""
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "售\n后"
    table.rows[0].cells[1].text = "七天"
    assert _table_to_lines(table) == ["售 后：七天"]


# ===== excel =====

def _write_xlsx(path, rows, columns=None, header=True):
    df = pd.DataFrame(rows, columns=columns)
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        df.to_excel(w, sheet_name="Sheet1", index=False, header=header)


def test_parse_excel_header_and_ffill(tmp_path):
    """有表头 sheet：输出 sheet 名 + 每行「列名：值」；合并单元格 NaN 被 ffill 补全。"""
    p = tmp_path / "a.xlsx"
    _write_xlsx(p, [["电池", "5800mAh"], [None, "120W"]], columns=["参数", "数值"])
    text = parse_excel(str(p))
    assert text == "【Sheet1】\n参数：电池；数值：5800mAh\n参数：电池；数值：120W"


def test_parse_excel_no_header_row_join(tmp_path):
    """首行为空的 sheet（列名 Unnamed）：退化为整行空格拼接。"""
    p = tmp_path / "b.xlsx"
    _write_xlsx(p, [[None, None], ["电池", "5800mAh"], ["快充", "120W"]], header=False)
    assert parse_excel(str(p)) == "电池 5800mAh\n快充 120W"


def test_parse_excel_empty_sheet_returns_empty(tmp_path):
    """仅表头无数据的 sheet 被跳过，返回空字符串。"""
    p = tmp_path / "c.xlsx"
    _write_xlsx(p, [], columns=["参数", "数值"])
    assert parse_excel(str(p)) == ""


def test_parse_excel_multi_sheets(tmp_path):
    """多 sheet 依次输出，各自带 sheet 名。"""
    p = tmp_path / "d.xlsx"
    with pd.ExcelWriter(p, engine="openpyxl") as w:
        pd.DataFrame([["电池", "5800mAh"]], columns=["参数", "数值"]).to_excel(w, sheet_name="手机", index=False)
        pd.DataFrame([["显卡", "RTX4060"]], columns=["参数", "数值"]).to_excel(w, sheet_name="笔记本", index=False)
    text = parse_excel(str(p))
    assert "【手机】" in text and "【笔记本】" in text
    assert "参数：显卡；数值：RTX4060" in text


# ===== pdf =====

def test_parse_pdf_primary_path_with_tables(monkeypatch):
    """pdfplumber 主链路：正文 + 表格行补充；全空表格行被过滤。"""
    import pdfplumber

    class _Page:
        def extract_text(self):
            return "电池容量 5800mAh"

        def extract_tables(self):
            return [[[None, ""], ["参数", "数值"], ["电池", "5800mAh"]]]

    class _Pdf:
        def __init__(self, pages):
            self.pages = pages

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    monkeypatch.setattr(pdfplumber, "open", lambda path: _Pdf([_Page()]))
    text = parse_pdf("fake.pdf")
    assert "电池容量 5800mAh" in text
    assert "参数；数值" in text
    assert "电池；5800mAh" in text


def test_parse_pdf_fallback_to_pypdf(monkeypatch):
    """pdfplumber 失败时回退 pypdf，页间空行拼接。"""
    import pdfplumber
    import pypdf

    class _FakePage:
        def __init__(self, text):
            self._text = text

        def extract_text(self):
            return self._text

    class _FakeReader:
        pages = [_FakePage("第一页内容"), _FakePage("第二页内容")]

        def __init__(self, path):
            pass

    def _boom(path):
        raise RuntimeError("pdfplumber 不可用")

    monkeypatch.setattr(pdfplumber, "open", _boom)
    monkeypatch.setattr(pypdf, "PdfReader", _FakeReader)
    assert parse_pdf("fake.pdf") == "第一页内容\n\n第二页内容"


# ===== 分发入口 =====

def test_parse_file_dispatch_happy_path(tmp_path):
    """按扩展名分发到对应 parser 并返回文本。"""
    p = tmp_path / "a.md"
    p.write_bytes("# 标题\n内容".encode("utf-8"))  # 显式字节避免 Windows 换行翻译
    assert parse_file(str(p), "md") == "# 标题\n内容"
    assert PARSERS["pdf"] is parse_pdf and PARSERS["xlsx"] is parse_excel


def test_parse_file_unsupported_type():
    """不支持的扩展名抛 ValueError。"""
    with pytest.raises(ValueError, match="不支持的文件类型"):
        parse_file("x.csv", "csv")


def test_parse_file_empty_result_raises(tmp_path):
    """解析结果为空（空白文件）抛 ValueError。"""
    p = tmp_path / "empty.txt"
    p.write_text("   ", encoding="utf-8")
    with pytest.raises(ValueError, match="解析结果为空"):
        parse_file(str(p), "txt")
